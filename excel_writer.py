from openpyxl import Workbook

SUMMARY_HEADERS = [
    "algorithm", "revenue", "total_cost", "revenuetocostratio",
    "accepted", "total_request", "embeddingratio", "pre_resource",
    "post_resource", "consumed", "avg_bw", "avg_crb", "avg_link",
    "No_of_Links_used", "avg_node", "No_of_Nodes_used", "avg_path",
    "avg_exec", "total_nodes", "total_links"
]

def write_results(filename, summary_rows, strategy_rows, accepted_rows,
                  resource_rows, pricing_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws.append(SUMMARY_HEADERS)
    for row in summary_rows:
        ws.append([row.get(h, "") for h in SUMMARY_HEADERS])

    ws2 = wb.create_sheet("VNR_Strategies")
    if strategy_rows:
        headers = list(strategy_rows[0].keys())
        ws2.append(headers)
        for row in strategy_rows:
            ws2.append([row.get(h, "") for h in headers])

    ws3 = wb.create_sheet("VNR_Results")
    if accepted_rows:
        headers = list(accepted_rows[0].keys())
        ws3.append(headers)
        for row in accepted_rows:
            ws3.append([row.get(h, "") for h in headers])

    ws4 = wb.create_sheet("Resources")
    if resource_rows:
        headers = list(resource_rows[0].keys())
        ws4.append(headers)
        for row in resource_rows:
            ws4.append([row.get(h, "") for h in headers])

    ws5 = wb.create_sheet("Upper_Level_DE")
    if pricing_rows:
        headers = list(pricing_rows[0].keys())
        ws5.append(headers)
        for row in pricing_rows:
            ws5.append([row.get(h, "") for h in headers])

    for wsx in wb.worksheets:
        wsx.freeze_panes = "A2"
        for cell in wsx[1]:
            cell.font = cell.font.copy(bold=True)
        for col in wsx.columns:
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            wsx.column_dimensions[col[0].column_letter].width = min(max_len + 2, 35)

    wb.save(filename)
